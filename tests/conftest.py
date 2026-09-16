"""Shared fixtures: a throwaway test database, and real PDF/DOCX builders.

Upload validation only checks that bytes look like a supported type, but the
ingestion pipeline actually parses them, so tests that exercise the full path
need genuine documents rather than byte stubs.
"""

import os
import re
import uuid
import zlib
from collections.abc import Callable, Iterator
from io import BytesIO

import psycopg
import pytest
from psycopg.conninfo import conninfo_to_dict, make_conninfo

# Keep app startup from running migrations under test; the `database` fixture
# below manages the schema itself.
os.environ.setdefault("APP_ENV", "test")
from docx import Document
from pypdf import PdfWriter

from app.api import dependencies
from app.database.migrations import run_migrations
from app.database.session import get_connection, get_database_url
from app.main import app
from app.retrieval.vector_store import VectorStore


# --- embeddings and vector store ---------------------------------------------


class FakeEmbedder:
    """Deterministic bag-of-words vectors: texts sharing words score higher.

    Enough to test indexing, filtering and ranking without loading a model.
    """

    backend = "fake"
    model_name = "fake-embedder"
    dimension = 64
    max_tokens = 512
    prompt_format = "none"

    def warm_up(self) -> None:
        pass

    def count_tokens(self, text: str) -> int:
        # One "token" per word: enough to exercise the chunker's token budget.
        return len(text.split())

    def _vector(self, text: str) -> list[float]:
        vector = [0.0] * self.dimension
        # crc32 rather than hash(): stable across processes, so tests can't
        # flake on a different collision pattern per run.
        for word in re.findall(r"[a-z0-9]+", text.lower()):
            vector[zlib.crc32(word.encode()) % self.dimension] += 1.0
        norm = sum(v * v for v in vector) ** 0.5 or 1.0
        return [v / norm for v in vector]

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._vector(t) for t in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._vector(text)


@pytest.fixture
def fake_embedder() -> FakeEmbedder:
    return FakeEmbedder()


class FakeChatModel:
    """Answers by citing the first passage, and records every prompt it was given.

    `reply` can be replaced per test (a fixed string, or a callable taking the
    user prompt) to script NOT_FOUND, uncited or malformed answers;
    `truncated` simulates running out of tokens.
    """

    provider = "fake"
    model_name = "fake-chat"

    def __init__(self) -> None:
        self.calls: list[tuple[str, str]] = []
        self.reply: str | Callable[[str], str] = "The passage states it [1]."
        self.truncated = False

    def complete(self, system: str, user: str, *, max_tokens: int):
        from app.answering.llm import Completion

        self.calls.append((system, user))
        return Completion(self.reply(user) if callable(self.reply) else self.reply, truncated=self.truncated)


@pytest.fixture
def fake_chat_model() -> FakeChatModel:
    return FakeChatModel()


@pytest.fixture
def vector_store(fake_embedder: FakeEmbedder) -> VectorStore:
    """An in-process Qdrant (no server) with the collection ready."""
    from qdrant_client import QdrantClient

    store = VectorStore(QdrantClient(":memory:"), collection="test_chunks")
    store.ensure_collection(fake_embedder.dimension)
    return store


@pytest.fixture(autouse=True)
def _fake_retrieval_stack(
    fake_embedder: FakeEmbedder, vector_store: VectorStore, fake_chat_model: FakeChatModel
) -> Iterator[None]:
    """Every API test gets the fake embedder, in-memory store and fake chat model by default."""
    app.dependency_overrides[dependencies.get_embedder] = lambda: fake_embedder
    app.dependency_overrides[dependencies.get_vector_store] = lambda: vector_store
    app.dependency_overrides[dependencies.get_chat_model] = lambda: fake_chat_model
    yield
    app.dependency_overrides.pop(dependencies.get_embedder, None)
    app.dependency_overrides.pop(dependencies.get_vector_store, None)
    app.dependency_overrides.pop(dependencies.get_chat_model, None)


# --- database ---------------------------------------------------------------


def _test_database_url(base_url: str) -> str:
    # Unique per session so concurrent runs (two agents, or a CI matrix) never
    # share — and truncate — the same tables.
    params = conninfo_to_dict(base_url)
    params["dbname"] = f"{params.get('dbname', 'postgres')}_test_{uuid.uuid4().hex[:8]}"
    return make_conninfo(**params)


def _create_database(base_url: str, name: str) -> None:
    # CREATE DATABASE cannot run inside a transaction, hence autocommit.
    with psycopg.connect(base_url, autocommit=True, connect_timeout=3) as connection:
        connection.execute(f'CREATE DATABASE "{name}"')


def _drop_database(base_url: str, name: str) -> None:
    with psycopg.connect(base_url, autocommit=True, connect_timeout=3) as connection:
        # FORCE closes any straggling connection so teardown can't leave the
        # database behind.
        connection.execute(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)')


@pytest.fixture(scope="session")
def database() -> Iterator[str]:
    """A migrated, session-unique `<dbname>_test_<id>` database.

    DATABASE_URL is pointed at it for the session and it is dropped afterwards.
    Skips when Postgres isn't reachable, unless MASIGN_REQUIRE_DB=1 (CI) makes
    that a failure instead. Tests never touch the dev database.
    """
    base_url = get_database_url()
    test_url = _test_database_url(base_url)
    name = conninfo_to_dict(test_url)["dbname"]
    try:
        _create_database(base_url, name)
    except psycopg.OperationalError as error:
        if os.getenv("MASIGN_REQUIRE_DB") == "1":
            raise
        pytest.skip(f"Postgres not reachable at {base_url}: {error}")

    previous = os.environ.get("DATABASE_URL")
    try:
        # Inside the try so a failing migration still drops the database.
        run_migrations(test_url)
        os.environ["DATABASE_URL"] = test_url
        yield test_url
    finally:
        if previous is None:
            os.environ.pop("DATABASE_URL", None)
        else:
            os.environ["DATABASE_URL"] = previous
        _drop_database(base_url, name)


@pytest.fixture
def db(database: str) -> Iterator[psycopg.Connection]:
    """A connection to the test database, emptied after each test."""
    with get_connection(database) as connection:
        yield connection
        connection.execute("TRUNCATE contracts, vector_index CASCADE")


# --- documents --------------------------------------------------------------


def build_pdf(lines: list[str]) -> bytes:
    """Build a minimal one-page PDF containing the given lines of text."""
    content = "BT /F1 12 Tf 14 TL 72 720 Td\n"
    for line in lines:
        escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        content += f"({escaped}) Tj T*\n"
    content += "ET"
    stream = content.encode("latin-1")

    page = (
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]"
        b"/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>"
    )
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        page,
        b"<</Length " + str(len(stream)).encode() + b">>stream\n" + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_position = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<</Size {len(objects) + 1}/Root 1 0 R>>\n"
        f"startxref\n{xref_position}\n%%EOF\n"
    ).encode()

    return bytes(out)


def build_scanned_pdf() -> bytes:
    """A structurally valid PDF with no text layer, like a scan."""
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def build_docx(
    paragraphs: list[str],
    table_rows: list[list[str]] | None = None,
) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)

    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for cell_index, value in enumerate(row):
                table.rows[row_index].cells[cell_index].text = value

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def make_pdf():
    return build_pdf


@pytest.fixture
def make_scanned_pdf():
    return build_scanned_pdf


@pytest.fixture
def make_docx():
    return build_docx
