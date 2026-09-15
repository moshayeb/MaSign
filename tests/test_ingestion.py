from itertools import pairwise
from pathlib import Path

import pytest

from app.ingestion.parsing import (
    DocumentParseError,
    NoExtractableTextError,
    UnsupportedFileTypeError,
    extract_text,
)
from app.ingestion.pipeline import TokenBudget, chunk_contract_text, load_contract_text


SAMPLE_CONTRACT = (
    Path(__file__).resolve().parent.parent
    / "data"
    / "sample_contracts"
    / "acme_vendor_agreement.txt"
)


# --- AA-22: extract raw text ------------------------------------------------


def test_extract_text_from_txt() -> None:
    content = b"1. Services\nVendor will provide services."

    assert extract_text(content, "txt") == "1. Services\nVendor will provide services."


def test_extract_text_drops_nul_characters() -> None:
    # MAS-42: extracted text must never carry NUL into the database.
    assert extract_text(b"1. Ser\x00vices\n\x00", "txt") == "1. Services"


def test_extract_text_strips_utf8_bom_and_crlf() -> None:
    content = "﻿ACME Vendor Agreement\r\n\r\n1. Services\r\n".encode("utf-8")

    assert extract_text(content, "txt") == "ACME Vendor Agreement\n\n1. Services"


def test_extract_text_from_pdf(make_pdf) -> None:
    pdf = make_pdf(["3. Limitation of Liability", "Vendor liability is capped."])

    text = extract_text(pdf, "pdf")

    assert "Limitation of Liability" in text
    assert "capped" in text


def test_extract_text_from_docx(make_docx) -> None:
    docx = make_docx(["4. Termination", "Either party may terminate for breach."])

    text = extract_text(docx, "docx")

    assert "4. Termination" in text
    assert "Either party may terminate for breach." in text


def test_extract_text_reads_docx_tables(make_docx) -> None:
    docx = make_docx(["2. Fees"], table_rows=[["Term", "36 months"]])

    text = extract_text(docx, "docx")

    assert "Term | 36 months" in text


def test_docx_tables_stay_in_document_order() -> None:
    # MAS-37: a table between two paragraphs must not be moved to the end.
    text = extract_text(_docx_with_table_in_middle(), "docx")

    assert text == "Before table\nTable term\nAfter table"


def test_docx_with_several_interleaved_tables_keeps_order() -> None:
    text = extract_text(_docx_interleaved(), "docx")

    assert text.split("\n") == ["P1", "T1", "P2", "T2 | T2b", "P3"]


def test_docx_nested_table_text_is_kept_in_order() -> None:
    # MAS-45: a table inside a cell used to vanish (cell.text ignores it).
    from io import BytesIO

    from docx import Document

    document = Document()
    document.add_paragraph("Fees")
    cell = document.add_table(rows=1, cols=1).rows[0].cells[0]
    cell.text = "Payment terms"
    cell.add_table(rows=1, cols=1).rows[0].cells[0].text = "Late payment penalty: 8 percent"
    document.add_paragraph("Term")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_text(buffer.getvalue(), "docx")

    assert text == "Fees\nPayment terms Late payment penalty: 8 percent\nTerm"


def test_docx_without_body_reports_parse_error() -> None:
    # MAS-46: well-formed XML with no <w:body> must be a parse error, not a crash.
    with pytest.raises(DocumentParseError, match="no document body"):
        extract_text(_docx_without_body(), "docx")


def _docx_without_body() -> bytes:
    from io import BytesIO
    from zipfile import ZipFile

    from docx import Document

    document = Document()
    document.add_paragraph("x")
    original = BytesIO()
    document.save(original)

    source = ZipFile(BytesIO(original.getvalue()))
    rebuilt = BytesIO()
    with ZipFile(rebuilt, "w") as archive:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                data = (
                    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"></w:document>'
                )
            archive.writestr(item, data)
    return rebuilt.getvalue()


def test_docx_with_malformed_xml_reports_parse_error() -> None:
    # MAS-36: a DOCX-shaped ZIP whose parts are not XML must be a clear parse
    # error, not an unhandled lxml exception.
    with pytest.raises(DocumentParseError):
        extract_text(_docx_with_broken_xml(), "docx")


def _docx_with_broken_xml() -> bytes:
    from io import BytesIO
    from zipfile import ZipFile

    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<broken")
        archive.writestr("word/document.xml", "<broken")
    return buffer.getvalue()


def _docx_with_table_in_middle() -> bytes:
    from io import BytesIO

    from docx import Document

    document = Document()
    document.add_paragraph("Before table")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Table term"
    document.add_paragraph("After table")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _docx_interleaved() -> bytes:
    from io import BytesIO

    from docx import Document

    document = Document()
    document.add_paragraph("P1")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "T1"
    document.add_paragraph("P2")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "T2"
    table.rows[0].cells[1].text = "T2b"
    document.add_paragraph("P3")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_scanned_pdf_reports_no_extractable_text(make_scanned_pdf) -> None:
    with pytest.raises(NoExtractableTextError):
        extract_text(make_scanned_pdf(), "pdf")


def test_corrupted_pdf_reports_parse_error() -> None:
    with pytest.raises(DocumentParseError):
        extract_text(b"%PDF-1.4\nnot actually a pdf body", "pdf")


def test_corrupted_docx_reports_parse_error() -> None:
    with pytest.raises(DocumentParseError):
        extract_text(b"PK\x03\x04 not a real docx package", "docx")


def test_empty_text_file_reports_no_extractable_text() -> None:
    with pytest.raises(NoExtractableTextError):
        extract_text(b"   \n\n  ", "txt")


def test_unknown_file_type_is_rejected() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        extract_text(b"anything", "rtf")


# --- AA-23: chunking --------------------------------------------------------


def test_chunking_keeps_short_contract_in_one_chunk() -> None:
    text = SAMPLE_CONTRACT.read_text(encoding="utf-8")

    chunks = chunk_contract_text(text)

    assert len(chunks) == 1
    assert "Limitation of Liability" in chunks[0]


def test_chunking_keeps_a_clause_heading_with_its_body() -> None:
    text = "1. Services\nVendor will provide services.\n\n2. Fees\nACME will pay."

    chunks = chunk_contract_text(text, max_chars=60, overlap_chars=0)

    assert chunks[0] == "1. Services\nVendor will provide services."
    assert chunks[1] == "2. Fees\nACME will pay."


def test_chunks_never_exceed_max_chars() -> None:
    text = "\n\n".join(f"{n}. Clause\n{'word ' * 60}".strip() for n in range(1, 12))

    chunks = chunk_contract_text(text, max_chars=400, overlap_chars=80)

    assert len(chunks) > 1
    assert all(len(chunk) <= 400 for chunk in chunks)


def test_consecutive_chunks_overlap() -> None:
    paragraphs = [f"{n}. Clause\n{'alpha beta gamma ' * 12}".strip() for n in range(1, 8)]

    chunks = chunk_contract_text("\n\n".join(paragraphs), max_chars=400, overlap_chars=100)

    assert len(chunks) > 1
    for previous, following in pairwise(chunks):
        carried = following.split("\n\n")[0]
        assert carried and carried in previous


def test_overlap_can_be_disabled() -> None:
    paragraphs = [f"{n}. Clause\n{'word ' * 40}".strip() for n in range(1, 6)]

    chunks = chunk_contract_text("\n\n".join(paragraphs), max_chars=300, overlap_chars=0)

    assert len(chunks) > 1
    assert sum(len(chunk) for chunk in chunks) < len("\n\n".join(paragraphs)) + len(chunks) * 4


def test_paragraph_longer_than_the_budget_is_split() -> None:
    text = "1. Services\n" + ("obligation " * 200).strip()

    chunks = chunk_contract_text(text, max_chars=300, overlap_chars=50)

    assert len(chunks) > 1
    assert all(len(chunk) <= 300 for chunk in chunks)
    assert "obligation" in chunks[-1]


def test_unbroken_run_of_characters_is_still_split() -> None:
    chunks = chunk_contract_text("x" * 500, max_chars=100, overlap_chars=0)

    assert len(chunks) == 5
    assert all(len(chunk) <= 100 for chunk in chunks)


def test_overlap_separator_is_counted_against_max_chars() -> None:
    # MAS-38: the "\n\n" between overlap and body used to push chunks over the limit.
    chunks = chunk_contract_text("x" * 250, max_chars=100, overlap_chars=20)

    assert all(len(chunk) <= 100 for chunk in chunks)
    assert all(chunk for chunk in chunks)
    # Nothing lost: every body piece is still present in order.
    assert "".join(chunk.split("\n\n")[-1] for chunk in chunks) == "x" * 250


def test_small_valid_budget_still_terminates_and_overlaps() -> None:
    chunks = chunk_contract_text("word " * 30, max_chars=20, overlap_chars=5)

    assert chunks
    assert all(0 < len(chunk) <= 20 for chunk in chunks)
    for previous, following in pairwise(chunks):
        carried = following.split("\n\n")[0]
        assert carried in previous


def test_overlap_that_leaves_no_room_for_a_body_is_rejected() -> None:
    with pytest.raises(ValueError):
        chunk_contract_text("text", max_chars=10, overlap_chars=8)  # 10 - 8 - 2 == 0


def test_blank_text_produces_no_chunks() -> None:
    assert chunk_contract_text("   \n\n  ") == []


def test_invalid_chunk_settings_are_rejected() -> None:
    with pytest.raises(ValueError):
        chunk_contract_text("text", max_chars=0)

    with pytest.raises(ValueError):
        chunk_contract_text("text", max_chars=100, overlap_chars=100)


# --- MAS-49: token budget ---------------------------------------------------


def _words(text: str) -> int:
    return len(text.split())


# Short words: 1200 characters hold far more than 40 of them, so only the
# token budget can be the reason to split.
DENSE_CLAUSE = " ".join(f"fee {n} due 1.5% cap 9,999" for n in range(30)) + " late penalty 8 percent"


def test_token_budget_splits_chunks_that_fit_by_characters() -> None:
    by_chars_only = chunk_contract_text(DENSE_CLAUSE, max_chars=1200, overlap_chars=0)
    assert len(by_chars_only) == 1  # the character limit alone would keep it whole

    chunks = chunk_contract_text(
        DENSE_CLAUSE, max_chars=1200, overlap_chars=0, token_budget=TokenBudget(_words, 40)
    )

    assert len(chunks) > 1
    assert all(_words(chunk) <= 40 for chunk in chunks)
    assert chunks[-1].endswith("late penalty 8 percent")  # the end of the clause survives
    assert " ".join(chunks).split() == DENSE_CLAUSE.split()  # nothing lost, nothing duplicated


def test_token_budget_counts_the_overlap_as_part_of_the_chunk() -> None:
    chunks = chunk_contract_text(
        DENSE_CLAUSE, max_chars=1200, overlap_chars=60, token_budget=TokenBudget(_words, 40)
    )

    assert len(chunks) > 1
    for previous, following in pairwise(chunks):
        carried = following.split("\n\n")[0]
        assert carried and previous.endswith(carried)  # the overlap really is there ...
    assert all(_words(chunk) <= 40 for chunk in chunks)  # ... and is counted in the budget


def test_token_budget_applies_when_packing_paragraphs_too() -> None:
    paragraphs = [f"{n}. Clause\n" + "word " * 15 for n in range(1, 6)]  # 17 words each

    chunks = chunk_contract_text(
        "\n\n".join(paragraphs), max_chars=1200, overlap_chars=0, token_budget=TokenBudget(_words, 40)
    )

    assert [chunk.count("Clause") for chunk in chunks] == [2, 2, 1]  # two clauses per 40-word chunk
    assert all(_words(chunk) <= 40 for chunk in chunks)


def _bodies(chunks: list[str]) -> list[str]:
    """Each chunk's own text, without the carried-over overlap."""
    return [chunk.split("\n\n")[-1] for chunk in chunks]


def test_overlap_shrinks_to_leave_the_body_half_the_token_budget() -> None:
    # MAS-57: 39-character, ten-word paragraphs; max_chars keeps one per chunk.
    # A 30-character overlap (~7 words) would eat most of a 12-word budget, so it
    # is cut down until it takes at most half — and nothing exceeds the limit.
    paragraphs = [" ".join(f"w{n}{i}" for i in range(10)) for n in range(4)]
    text = "\n\n".join(paragraphs)

    chunks = chunk_contract_text(text, max_chars=75, overlap_chars=30, token_budget=TokenBudget(_words, 12))

    assert all(_words(chunk) <= 12 for chunk in chunks)
    assert " ".join(_bodies(chunks)).split() == text.split()  # nothing lost or duplicated
    assert any("\n\n" in chunk for chunk in chunks)  # a (shorter) overlap survives
    assert all(_words(chunk.split("\n\n")[0]) <= 6 for chunk in chunks if "\n\n" in chunk)


def test_overlap_is_dropped_when_even_a_sliver_cannot_fit() -> None:
    # MAS-57: a model whose prefix and special tokens (9) already exceed half of
    # a 16-token budget leaves no room for any overlap; the clauses stay intact.
    paragraphs = [" ".join(f"w{n}{i}" for i in range(6)) for n in range(4)]  # 23 chars each

    chunks = chunk_contract_text(
        "\n\n".join(paragraphs), max_chars=45, overlap_chars=20, token_budget=TokenBudget(lambda t: 9 + _words(t), 16)
    )

    assert chunks == paragraphs


def test_overlap_is_kept_where_the_token_budget_allows_it() -> None:
    paragraphs = [" ".join(f"w{n}{i}" for i in range(10)) for n in range(3)]

    chunks = chunk_contract_text(
        "\n\n".join(paragraphs), max_chars=75, overlap_chars=30, token_budget=TokenBudget(_words, 40)
    )

    assert chunks[0] == paragraphs[0]
    assert chunks[1].endswith(paragraphs[1]) and chunks[1] != paragraphs[1]  # overlap present


def test_token_budget_too_small_for_any_chunk_is_rejected() -> None:
    # A budget the prefix and special tokens alone exhaust can never be met.
    with pytest.raises(ValueError, match="token budget"):
        chunk_contract_text("some text", token_budget=TokenBudget(lambda t: 10 + _words(t), 8))


def test_token_budget_falls_back_to_character_cuts_for_unbroken_runs() -> None:
    # A single 300-character "word" can never fit a budget of 100 characters;
    # the token measure (one token per 10 characters) must not loop forever.
    chunks = chunk_contract_text("x" * 300, max_chars=100, overlap_chars=0, token_budget=TokenBudget(lambda t: len(t) // 10, 5))

    assert chunks
    assert all(len(chunk) // 10 <= 5 for chunk in chunks)
    assert "".join(chunks) == "x" * 300


# --- end to end -------------------------------------------------------------


def test_sample_contract_loads_and_chunks() -> None:
    text = load_contract_text(SAMPLE_CONTRACT)

    chunks = chunk_contract_text(text, max_chars=120, overlap_chars=20)

    assert len(chunks) > 1
    assert all(chunk.strip() for chunk in chunks)

    # Splitting happens on word boundaries, so every word of the contract
    # survives somewhere even when a clause is broken across chunks.
    joined = " ".join(chunks)
    for word in text.split():
        assert word in joined
