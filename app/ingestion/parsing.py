"""Text extraction for the supported contract file types.

Upload validation (AA-45) decides a file's type from its bytes and hands on a
``file_type`` of ``txt``, ``pdf`` or ``docx``. This module turns those bytes
into plain text. It deliberately imports nothing from the web layer, so the
same functions work from scripts, tests and the API.
"""

import re
from dataclasses import dataclass, field
from io import BytesIO
from zipfile import BadZipFile

import pypdf
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml.etree import XMLSyntaxError


SUPPORTED_FILE_TYPES = ("txt", "pdf", "docx")


@dataclass(frozen=True)
class ExtractedDocument:
    """The text of an upload plus what could not be read from it (MAS-84).

    `notes` are plain sentences for the user ("Pages 3 and 7 have no text
    layer …"); they are stored with the contract and shown as coverage, so
    that a review of the readable part is never mistaken for a review of the
    whole document.
    """

    text: str
    notes: list[str] = field(default_factory=list)


class DocumentTextError(ValueError):
    """Base class for failures to get text out of an uploaded document."""


class UnsupportedFileTypeError(DocumentTextError):
    """Raised for a file type the ingestion pipeline does not handle."""


class DocumentParseError(DocumentTextError):
    """Raised when a document cannot be opened or read at all."""


class NoExtractableTextError(DocumentTextError):
    """Raised when a document opens cleanly but holds no usable text.

    The case that matters in practice is a scanned contract: every page is an
    image, so the file is a structurally valid PDF and passes upload validation,
    but there is no text layer to read. Callers turn this into a clear
    unsupported-document message rather than an empty result.
    """


def extract_text(content: bytes, file_type: str) -> str:
    """The plain text of a validated upload (see `extract_document` for what was not read)."""
    return extract_document(content, file_type).text


def extract_document(content: bytes, file_type: str) -> ExtractedDocument:
    """Return the plain text of a validated upload and notes on what could not be read.

    Raises `UnsupportedFileTypeError` for unknown types, `DocumentParseError`
    if the bytes cannot be read, and `NoExtractableTextError` if they can be
    read but contain no text.
    """
    notes: list[str] = []
    if file_type == "txt":
        text = _extract_txt(content)
    elif file_type == "pdf":
        text = _extract_pdf(content, notes)
    elif file_type == "docx":
        text = _extract_docx(content)
    else:
        raise UnsupportedFileTypeError(
            f"Cannot extract text from file type {file_type!r}. "
            f"Supported types are {', '.join(SUPPORTED_FILE_TYPES)}."
        )

    removed = text.count("\x00")
    if removed:
        notes.append(f"{removed} unreadable character{'' if removed == 1 else 's'} removed from the text.")
    text = normalize_text(text)
    if not text:
        raise NoExtractableTextError(
            "No readable text found in the document. Scanned contracts must be "
            "run through OCR before they can be processed."
        )

    return ExtractedDocument(text, notes)


def normalize_text(text: str) -> str:
    """Normalise line endings and whitespace without losing paragraph breaks."""
    # PDF and DOCX extraction can surface NUL characters from odd encodings;
    # they carry no text and PostgreSQL text columns reject them.
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    # Collapse runs of blank lines so paragraph splitting stays predictable.
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_txt(content: bytes) -> str:
    # Validation already confirmed this decodes as UTF-8; utf-8-sig drops a BOM
    # if the file came from a Windows editor.
    return content.decode("utf-8-sig", errors="replace")


def _extract_pdf(content: bytes, notes: list[str] | None = None) -> str:
    try:
        reader = pypdf.PdfReader(BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as error:  # pypdf raises a wide range of parse errors
        raise DocumentParseError(
            "The PDF could not be read. It may be corrupted or password protected."
        ) from error

    # A page with no text layer (a scan, a drawing, a signature page) is
    # not read at all; the contract must say so rather than look complete.
    empty = [number for number, page in enumerate(pages, start=1) if not page.strip()]
    if notes is not None and empty and len(empty) < len(pages):
        notes.append(
            f"{_page_list(empty)} of {len(pages)} {'has' if len(empty) == 1 else 'have'} no text layer "
            "(scanned or image-only) and could not be read."
        )
    return "\n\n".join(page for page in pages if page.strip())


def _page_list(numbers: list[int]) -> str:
    if len(numbers) == 1:
        return f"Page {numbers[0]}"
    if len(numbers) <= 6:
        return "Pages " + ", ".join(str(n) for n in numbers[:-1]) + f" and {numbers[-1]}"
    return f"{len(numbers)} pages"


def _extract_docx(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))
    except (PackageNotFoundError, BadZipFile, KeyError, ValueError, XMLSyntaxError) as error:
        # XMLSyntaxError: a ZIP shaped like a DOCX whose parts are not valid XML.
        raise DocumentParseError(
            "The DOCX file could not be read. It may be corrupted."
        ) from error

    body = document.element.body
    if body is None:
        # Well-formed XML, but no <w:body>: nothing python-docx can read from.
        raise DocumentParseError("The DOCX file could not be read. It has no document body.")

    return "\n".join(_iter_block_text(body, document))


def _iter_block_text(container, document) -> list[str]:
    """Paragraph and table text in document order, recursing into table cells.

    Walking the XML children (rather than document.paragraphs + document.tables)
    keeps a table between the paragraphs that surround it, so fees, dates or
    parties stay attached to their clause. Cells are walked the same way, so a
    table nested inside a cell is not lost.
    """
    blocks: list[str] = []
    for element in container.iterchildren():
        if element.tag == qn("w:p"):
            blocks.append(Paragraph(element, document).text)
        elif element.tag == qn("w:tbl"):
            for row in Table(element, document).rows:
                cells = [" ".join(_iter_block_text(cell._tc, document)).strip() for cell in row.cells]
                line = " | ".join(cell for cell in cells if cell)
                if line:
                    blocks.append(line)
    return blocks
